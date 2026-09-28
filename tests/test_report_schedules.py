"""Scheduled report delivery: tests written before the code.

Contract under test (api/report_schedules.py + api/report_schedule_routes.py):
  - A schedule subscribes recipients to a SAVED VIEW on a cadence (daily/weekly/
    monthly) with a trailing data window; creation validates the view exists in
    the org, the recipients are real addresses, and the cadence fields are sane.
  - is_due() is pure calendar logic: due once per period at/after the configured
    UTC hour, never twice in the same period.
  - rows_to_csv() renders business labels as the header, booleans as True/False,
    and None as empty — same conventions as the SPA's export.
  - run_due_schedules() renders each due schedule and hands an email (with the
    CSV attached) to a send callable; dry-run renders but never sends; a failing
    schedule records its error and does not block the rest.
  - Routes are org-scoped: a user manages only their org's schedules.
"""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

# Auth-disabled app context (same pattern as the other route tests): the API acts
# as the dev org without tokens.
os.environ["PORTAL_AUTH_DISABLED"] = "true"
os.environ["PORTAL_DEV_ORGANIZATION"] = "dev"
os.environ.pop("PORTAL_STATE_DATABASE_URL", None)
os.environ.pop("PORTAL_AUTH_DATABASE_URL", None)

from api import report_schedules as rs  # noqa: E402

UTC = timezone.utc
VIEW = {
    "id": "view-1",
    "organization_id": "dev",
    "snapshot_id": "rpt_bill_segment",
    "snapshot_label": "Bill Segment",
    "title": "Frozen segments by division",
    "kind": "builder",
    "dimensions": ["division"],
    "measures": [{"field": "*", "agg": "count"}],
}


def _payload(**over):
    base = {
        "saved_view_id": "view-1",
        "recipients": ["exec@utility.gov"],
        "cadence": "daily",
        "hour_utc": 13,
        "window_days": 30,
    }
    base.update(over)
    return base


class ScheduleStoreTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._path = mock.patch.object(
            rs, "SCHEDULES_PATH", Path(self._tmp.name) / "report_schedules.json")
        self._path.start()
        self._views = mock.patch.object(rs, "_find_view", side_effect=lambda vid, org: (
            VIEW if (vid == "view-1" and org == "dev") else None))
        self._views.start()

    def tearDown(self):
        self._views.stop()
        self._path.stop()
        self._tmp.cleanup()

    def test_create_and_list(self):
        entry = rs.create_schedule(_payload(), organization_id="dev", created_by="a@b.gov")
        self.assertTrue(entry["id"])
        self.assertEqual(entry["saved_view_id"], "view-1")
        self.assertEqual(entry["created_by"], "a@b.gov")
        listed = rs.list_schedules("dev")
        self.assertEqual([s["id"] for s in listed], [entry["id"]])
        # other org sees nothing
        self.assertEqual(rs.list_schedules("ellensburg"), [])

    def test_create_rejects_unknown_view(self):
        with self.assertRaises(rs.ScheduleError):
            rs.create_schedule(_payload(saved_view_id="nope"),
                               organization_id="dev", created_by="a@b.gov")

    def test_create_rejects_bad_recipient_and_cadence(self):
        with self.assertRaises(rs.ScheduleError):
            rs.create_schedule(_payload(recipients=["not-an-email"]),
                               organization_id="dev", created_by="a@b.gov")
        with self.assertRaises(rs.ScheduleError):
            rs.create_schedule(_payload(recipients=[]),
                               organization_id="dev", created_by="a@b.gov")
        with self.assertRaises(rs.ScheduleError):
            rs.create_schedule(_payload(cadence="hourly"),
                               organization_id="dev", created_by="a@b.gov")
        with self.assertRaises(rs.ScheduleError):
            rs.create_schedule(_payload(cadence="weekly", weekday=9),
                               organization_id="dev", created_by="a@b.gov")

    def test_delete_is_org_scoped(self):
        entry = rs.create_schedule(_payload(), organization_id="dev", created_by="a@b.gov")
        self.assertFalse(rs.delete_schedule(entry["id"], "ellensburg"))
        self.assertTrue(rs.delete_schedule(entry["id"], "dev"))
        self.assertEqual(rs.list_schedules("dev"), [])


class DueLogicTests(unittest.TestCase):
    def _sched(self, **over):
        base = {"cadence": "daily", "hour_utc": 13, "enabled": True, "last_run_at": None}
        base.update(over)
        return base

    def test_daily_due_at_hour_once(self):
        s = self._sched()
        self.assertFalse(rs.is_due(s, datetime(2026, 9, 1, 12, 59, tzinfo=UTC)))
        self.assertTrue(rs.is_due(s, datetime(2026, 9, 1, 13, 0, tzinfo=UTC)))
        s["last_run_at"] = datetime(2026, 9, 1, 13, 1, tzinfo=UTC).isoformat()
        self.assertFalse(rs.is_due(s, datetime(2026, 9, 1, 18, 0, tzinfo=UTC)))
        self.assertTrue(rs.is_due(s, datetime(2026, 9, 2, 13, 0, tzinfo=UTC)))

    def test_weekly_due_on_weekday(self):
        s = self._sched(cadence="weekly", weekday=0)  # Monday
        self.assertTrue(rs.is_due(s, datetime(2026, 9, 7, 13, 0, tzinfo=UTC)))   # a Monday
        self.assertFalse(rs.is_due(s, datetime(2026, 9, 8, 13, 0, tzinfo=UTC)))  # Tuesday
        s["last_run_at"] = datetime(2026, 9, 7, 13, 1, tzinfo=UTC).isoformat()
        self.assertFalse(rs.is_due(s, datetime(2026, 9, 7, 20, 0, tzinfo=UTC)))

    def test_monthly_due_first_of_month(self):
        s = self._sched(cadence="monthly")
        self.assertTrue(rs.is_due(s, datetime(2026, 9, 1, 13, 0, tzinfo=UTC)))
        self.assertFalse(rs.is_due(s, datetime(2026, 9, 2, 13, 0, tzinfo=UTC)))

    def test_disabled_never_due(self):
        s = self._sched(enabled=False)
        self.assertFalse(rs.is_due(s, datetime(2026, 9, 1, 13, 0, tzinfo=UTC)))


class CsvRenderTests(unittest.TestCase):
    def test_labels_booleans_and_none(self):
        csv_text = rs.rows_to_csv(
            columns=["division", "is_frozen", "amt"],
            labels={"division": "Division", "is_frozen": "Is Frozen", "amt": "Amount"},
            rows=[
                {"division": "Water", "is_frozen": True, "amt": 12.5},
                {"division": "Electric", "is_frozen": False, "amt": None},
            ])
        lines = csv_text.strip().splitlines()
        self.assertEqual(lines[0], "Division,Is Frozen,Amount")
        self.assertEqual(lines[1], "Water,True,12.5")
        self.assertEqual(lines[2], "Electric,False,")


class RunnerTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._path = mock.patch.object(
            rs, "SCHEDULES_PATH", Path(self._tmp.name) / "report_schedules.json")
        self._path.start()
        self._views = mock.patch.object(rs, "_find_view", side_effect=lambda vid, org: (
            VIEW if vid == "view-1" else None))
        self._views.start()
        self._create("a@b.gov")

    def _create(self, who):
        # The runs below are dated 1 Sep 2026; a schedule is never owed a moment from
        # before it existed, so the fixture exists from August.
        entry = rs.create_schedule(_payload(), organization_id="dev", created_by=who)
        entry["created_at"] = datetime(2026, 8, 1, tzinfo=UTC).isoformat()
        rs._store.update(entry)

    def tearDown(self):
        self._views.stop()
        self._path.stop()
        self._tmp.cleanup()

    def _render_ok(self, schedule, view):
        return (["division"], {"division": "Division"}, [{"division": "Water"}])

    def test_run_sends_email_with_csv_attachment(self):
        sent = []
        now = datetime(2026, 9, 1, 13, 5, tzinfo=UTC)
        with mock.patch.object(rs, "render_schedule", side_effect=self._render_ok):
            results = rs.run_due_schedules(now=now, send=lambda msg: sent.append(msg))
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["status"], "sent")
        self.assertEqual(len(sent), 1)
        msg = sent[0]
        self.assertIn("exec@utility.gov", msg["To"])
        self.assertIn("Frozen segments by division", msg["Subject"])
        attachments = [p for p in msg.iter_attachments()]
        self.assertEqual(len(attachments), 1)
        body = attachments[0].get_content()
        self.assertIn("Division", body)
        self.assertIn("Water", body)
        # marked run: not due again today
        again = rs.run_due_schedules(now=now, send=lambda msg: sent.append(msg))
        self.assertEqual(again, [])

    def test_dry_run_never_sends(self):
        sent = []
        now = datetime(2026, 9, 1, 13, 5, tzinfo=UTC)
        with mock.patch.object(rs, "render_schedule", side_effect=self._render_ok):
            results = rs.run_due_schedules(now=now, send=lambda m: sent.append(m), dry_run=True)
        self.assertEqual(results[0]["status"], "dry-run")
        self.assertEqual(sent, [])
        # dry-run must NOT mark the schedule as run
        with mock.patch.object(rs, "render_schedule", side_effect=self._render_ok):
            live = rs.run_due_schedules(now=now, send=lambda m: sent.append(m))
        self.assertEqual(live[0]["status"], "sent")

    def test_render_failure_recorded_not_fatal(self):
        self._create("b@b.gov")
        now = datetime(2026, 9, 1, 13, 5, tzinfo=UTC)
        calls = {"n": 0}

        def flaky(schedule, view):
            calls["n"] += 1
            if calls["n"] == 1:
                raise RuntimeError("warehouse down")
            return self._render_ok(schedule, view)

        sent = []
        with mock.patch.object(rs, "render_schedule", side_effect=flaky):
            results = rs.run_due_schedules(now=now, send=lambda m: sent.append(m))
        statuses = sorted(r["status"].split(":")[0] for r in results)
        self.assertEqual(statuses, ["error", "sent"])
        self.assertEqual(len(sent), 1)


class RouteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from api.report_schedule_routes import router

        cls._tmp = tempfile.TemporaryDirectory()
        cls._path = mock.patch.object(
            rs, "SCHEDULES_PATH", Path(cls._tmp.name) / "report_schedules.json")
        cls._path.start()
        cls._views = mock.patch.object(rs, "_find_view", side_effect=lambda vid, org: (
            VIEW if (vid == "view-1" and org == "dev") else None))
        cls._views.start()
        # another module may have flipped auth on at import time; pin it per-class
        # (auth_disabled() reads the env per request)
        cls._env = mock.patch.dict(os.environ, {
            "PORTAL_AUTH_DISABLED": "true", "PORTAL_DEV_ORGANIZATION": "dev",
            # dev needs a data source: require_org_for_data no longer accepts
            # the global credential fallback (audit H2).
            "WAREHOUSE_DATABASE_URL": "postgresql://test@localhost/test",
        })
        cls._env.start()
        app = FastAPI()
        app.include_router(router)
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        cls._env.stop()
        cls._views.stop()
        cls._path.stop()
        cls._tmp.cleanup()

    def test_crud_roundtrip(self):
        r = self.client.post("/report-schedules", json=_payload())
        self.assertEqual(r.status_code, 200, r.text)
        sid = r.json()["id"]
        r = self.client.get("/report-schedules")
        self.assertIn(sid, [s["id"] for s in r.json()["schedules"]])
        r = self.client.delete(f"/report-schedules/{sid}")
        self.assertEqual(r.status_code, 200)
        r = self.client.get("/report-schedules")
        self.assertEqual(r.json()["schedules"], [])

    def test_create_bad_payload_is_400(self):
        r = self.client.post("/report-schedules", json=_payload(recipients=["nope"]))
        self.assertEqual(r.status_code, 400)


if __name__ == "__main__":
    unittest.main()


class ScheduleIsFaithfulTests(unittest.TestCase):
    """A scheduled report must be the view the person saved, delivered when it was due.

    Found by the parity inventory (2026-09-28): the runner applied the date window and
    the legacy single scope pair but never the view's saved `filters`, so the emailed
    CSV could differ from the view on screen; and a monthly schedule fired only when the
    runner happened to run on the 1st, so one missed hour lost the whole month."""

    def _render(self, view):
        seen = {}

        def build(**kw):
            seen.update(kw)
            return "select 1", {}
        snap = {"table_name": "rpt_bill_segment", "trusted_measures": []}
        with mock.patch("api.snapshot_catalog.get_snapshot", return_value=snap), \
             mock.patch("api.snapshot_catalog.snapshot_backend", return_value=("postgres", "postgres", "reporting")), \
             mock.patch("api.snapshot_catalog.allowed_fields", return_value=set()), \
             mock.patch.object(rs, "schedule_date_field", return_value="Bill Date"), \
             mock.patch("api.reporting_dates.reporting_window", return_value=("2026-05-20", "2026-06-18")), \
             mock.patch("api.query_builder.build_query", side_effect=build), \
             mock.patch("api.warehouse_db.execute_query", return_value=(["m0"], [])):
            rs.render_schedule({"organization_id": "dev", "window_days": 30}, view)
        return seen["filters"]

    def test_the_saved_filters_are_applied(self):
        filters = self._render({"snapshot_id": "rpt_bill_segment", "dimensions": ["Bill Cycle"],
                                "measures": [{"field": "Billed Amount", "agg": "sum"}],
                                "filters": [{"field": "Is Frozen", "op": "eq", "value": True}]})
        self.assertIn({"field": "Is Frozen", "op": "eq", "value": True}, filters)

    def test_a_saved_date_range_gives_way_to_the_schedules_trailing_window(self):
        filters = self._render({"snapshot_id": "rpt_bill_segment",
                                "filters": [{"field": "Bill Date", "op": "between", "value": ["2025-01-01", "2025-01-31"]}]})
        dated = [f for f in filters if f["field"] == "Bill Date"]
        self.assertEqual(dated, [{"field": "Bill Date", "op": "between", "value": ["2026-05-20", "2026-06-18"]}])

    def test_a_filter_asked_for_on_open_runs_on_its_saved_value(self):
        """Nobody is there to answer a report parameter at 06:00; the schedule uses the
        value the view was saved with."""
        filters = self._render({"snapshot_id": "rpt_bill_segment",
                                "filters": [{"field": "Bill Cycle", "op": "eq", "value": "C1", "prompt": True}]})
        self.assertIn("C1", [f.get("value") for f in filters if f["field"] == "Bill Cycle"])

    def test_a_missed_monthly_run_is_caught_up_once(self):
        s = {"cadence": "monthly", "hour_utc": 13, "enabled": True,
             "last_run_at": datetime(2026, 8, 1, 13, 5, tzinfo=UTC).isoformat()}
        self.assertTrue(rs.is_due(s, datetime(2026, 9, 3, 9, 0, tzinfo=UTC)), "the runner was down on the 1st")
        s["last_run_at"] = datetime(2026, 9, 3, 9, 1, tzinfo=UTC).isoformat()
        self.assertFalse(rs.is_due(s, datetime(2026, 9, 4, 13, 0, tzinfo=UTC)))

    def test_a_new_schedule_waits_for_its_first_scheduled_moment(self):
        s = {"cadence": "monthly", "hour_utc": 13, "enabled": True, "last_run_at": None,
             "created_at": datetime(2026, 9, 10, 8, 0, tzinfo=UTC).isoformat()}
        self.assertFalse(rs.is_due(s, datetime(2026, 9, 12, 13, 0, tzinfo=UTC)))
        self.assertTrue(rs.is_due(s, datetime(2026, 10, 1, 13, 0, tzinfo=UTC)))


class ExcelAttachmentTests(unittest.TestCase):
    """A schedule can deliver Excel as well as CSV: numbers stay numbers, the header is the
    business labels, and the file opens as a real workbook."""

    def test_rows_become_a_real_workbook(self):
        import io
        from openpyxl import load_workbook
        data = rs.rows_to_xlsx(["cycle", "amt", "frozen"], {"cycle": "Bill Cycle", "amt": "Billed Amount", "frozen": "Is Frozen"},
                               [{"cycle": "C1", "amt": 2682879.14, "frozen": True}, {"cycle": "C2", "amt": None, "frozen": False}])
        ws = load_workbook(io.BytesIO(data)).active
        self.assertEqual([c.value for c in ws[1]], ["Bill Cycle", "Billed Amount", "Is Frozen"])
        self.assertEqual(ws["B2"].value, 2682879.14)
        self.assertIsNone(ws["B3"].value)
        self.assertEqual(ws["C2"].value, "True")

    def test_an_excel_schedule_attaches_xlsx(self):
        msg = rs._message({"view_title": "Billed by cycle", "recipients": ["a@b.gov"], "format": "xlsx"},
                          ["cycle"], {"cycle": "Bill Cycle"}, [{"cycle": "C1"}], datetime(2026, 9, 1, tzinfo=UTC))
        (att,) = list(msg.iter_attachments())
        self.assertTrue(att.get_filename().endswith(".xlsx"))
        self.assertIn("spreadsheetml", att.get_content_type())

    def test_an_unknown_format_is_refused(self):
        with self.assertRaisesRegex(rs.ScheduleError, "[Ff]ormat"):
            rs.create_schedule({**_payload(), "format": "docx"}, organization_id="dev", created_by="a@b.gov")


class PdfAttachmentTests(unittest.TestCase):
    """A schedule can deliver a formatted PDF built on the server: the Origin mark, title,
    the data window it applied, a table whose header repeats on every page, and page
    numbers. Every PDF in the portal before this was the browser's print dialog, which
    cannot run on a schedule."""

    ROWS = [{"cycle": f"Cycle {i}", "amt": 1000.5 + i} for i in range(120)]

    def _pdf(self, rows):
        return rs.rows_to_pdf("Billed by cycle", "Data window: trailing 30 days on Bill Date, as of 2026-06-18.",
                              ["cycle", "amt"], {"cycle": "Bill Cycle", "amt": "Billed Amount"}, rows,
                              datetime(2026, 9, 28, tzinfo=UTC))

    def test_it_is_a_pdf(self):
        self.assertTrue(self._pdf(self.ROWS[:3]).startswith(b"%PDF-"))

    def test_a_long_table_runs_to_several_pages(self):
        pages = self._pdf(self.ROWS).count(b"/Type /Page\n") + self._pdf(self.ROWS).count(b"/Type /Page ")
        self.assertGreater(pages, 1)

    def test_a_pdf_schedule_attaches_pdf(self):
        msg = rs._message({"view_title": "Billed by cycle", "recipients": ["a@b.gov"], "format": "pdf"},
                          ["cycle"], {"cycle": "Bill Cycle"}, [{"cycle": "C1"}], datetime(2026, 9, 1, tzinfo=UTC))
        (att,) = list(msg.iter_attachments())
        self.assertEqual(att.get_content_type(), "application/pdf")
        self.assertTrue(att.get_filename().endswith(".pdf"))
